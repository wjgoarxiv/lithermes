#!/usr/bin/env python3
"""docx_gate.py — the LitHermes pass/fail gate for one generated Word document.

Run it through the shared Office runner so the pinned runtime is used:

    node <installed-plugin>/skills/lit-pptx/bin/office.mjs gate report.docx --source report.md --layout
    node <installed-plugin>/skills/lit-pptx/bin/office.mjs gate report.docx --source report.md --compare other.docx
    node <installed-plugin>/skills/lit-pptx/bin/office.mjs gate paper.docx --source paper.md --publisher elsevier --kind manuscript

It reads, in order:

  package   the archive opens, the core parts are present and python-docx reopens it
  content   body text exists, every Markdown heading reached the page, no unfilled blank
            ([금액], XXX, TBD, ○○, {{…}}) and no frontmatter printed as body text
  prose     slop_lint over the Markdown (a report drops the manuscript-outline rules)
  design    slop_lint's DOCX design audit, and the Hangul font pairing on Korean text
  output    docx_layout: labels instead of sentences in the title and headings, heading order, the
            restraint checks (one accent, ink headings, quiet heading ratio, unfilled tables, the
            component budget, no ISO dates in Korean text, no coloured header chip); with --layout the
            document is rendered by LibreOffice under a private profile and its pages are checked for
            fill, stranded headings, word-safe heading wraps, splits and component variety; with
            --compare the two builds must differ in at least three structural features

Prints JSON; exit 0 on PASS, 1 on FAIL. A pass means no defect was found, not that the pages read
well: look at the rendered pages before calling the document done.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

sys.dont_write_bytecode = True
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from docx import Document  # noqa: E402

import docx_layout  # noqa: E402
import slop_lint  # noqa: E402
from convert_md_to_docx import parse_frontmatter  # noqa: E402

HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
# A bracket holding a digit is a citation ([1], [S3]); a blank holds none.
BLANK = re.compile(r"\[(?![^\]\n]*\d)[^\]\n]{1,24}\]|\bX{3,}\b|\bTBD\b|\bTODO\b|○○|OOO|\{\{[^}]*\}\}|_{4,}|＿{2,}")
FRONTMATTER_LINE = re.compile(r"^(?:---|title:|author:|authors:|date:|subtitle:)\s*", re.M)
SHARED_DESIGN_RULES = {"rule-45-highlighted-text", "rule-46-tracked-changes", "rule-55-ellipsis", "rule-60-numeric-column-alignment"}


def package_problems(path: Path) -> list[str]:
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if archive.testzip() is not None:
                return ["damaged archive member"]
    except (zipfile.BadZipFile, OSError) as error:
        return [f"not a readable package: {error}"]
    missing = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"} - names
    if missing:
        return [f"missing part {name}" for name in sorted(missing)]
    try:
        Document(str(path))
    except Exception as error:  # python-docx raises several unrelated types on a broken part
        return [f"python-docx cannot reopen it: {error}"]
    return []


def markdown_headings(text: str) -> list[str]:
    found, fenced = [], False
    for line in text.splitlines():
        if line.lstrip().startswith(("```", "~~~")):
            fenced = not fenced
            continue
        match = None if fenced else HEADING.match(line)
        if match:
            found.append(re.sub(r"[*_`]", "", match.group(2)).strip())
    return found


def content_report(path: Path, source_text: str | None) -> dict:
    document = Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    blanks = sorted({m.group(0) for text in paragraphs + cells for m in BLANK.finditer(text)})
    leaked = [p for p in paragraphs[:8] if FRONTMATTER_LINE.match(p)]
    missing = []
    if source_text:
        page_text = docx_layout.norm("".join(paragraphs + cells))  # a heading may sit inside a side-by-side row
        missing = [h for h in markdown_headings(source_text) if docx_layout.norm(h) not in page_text]
    chars = sum(len(p) for p in paragraphs) + sum(len(c) for c in cells)
    return {"pass": chars > 0 and not missing and not blanks and not leaked, "body_chars": chars,
            "placeholders": blanks[:20], "frontmatter_leak": leaked[:4], "missing_headings": missing[:20]}


def lint_report(source: Path | None, path: Path, publisher_name: str, kind: str, named_publisher: bool) -> dict:
    registry = HERE.parent / "templates" / "registry.yaml"
    phrases = slop_lint.load_phrase_rules(HERE.parent / "references" / "slop_phrase_list.yaml")
    publisher = slop_lint.load_publisher(registry, publisher_name)
    prose, locale = [], "auto"
    if source is not None:
        prose, _front, _body, locale = slop_lint.lint_text(source.read_text(encoding="utf-8"), publisher, phrases, "auto")
        if kind == "report":
            prose = [f for f in prose if f.rule_id not in slop_lint.MANUSCRIPT_ONLY_RULES]
        elif not named_publisher:
            prose = [f for f in prose if f.rule_id != "rule-08-structure-order"]  # no publisher outline to follow
    design = [f for f in slop_lint.audit_docx_design(path, publisher) if getattr(f, "severity", "HIGH") == "HIGH"]
    if locale in {"ko", "mixed"} or source is None:
        design += slop_lint.audit_docx_cjk(path)
    if not named_publisher:
        # The plain and tonality builds have no registry profile: only the rules every profile shares apply.
        design = [f for f in design if f.rule_id in SHARED_DESIGN_RULES]
    row = lambda f: {"rule": f.rule_id, "line": f.line, "section": f.section, "message": f.message, "excerpt": f.excerpt}  # noqa: E731
    return {"pass": not prose and not design, "locale": locale,
            "prose": [row(f) for f in prose][:40], "design": [row(f) for f in design][:40]}


def render_pdf(path: Path, out_dir: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory(prefix="lithermes-soffice-") as profile:
        subprocess.run([soffice, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(out_dir), str(path)], capture_output=True, text=True, timeout=300)
    pdf = out_dir / f"{path.stem}.pdf"
    return pdf if pdf.exists() else None


def output_report(path: Path, source: Path | None, args) -> dict:
    front = parse_frontmatter(source.read_text(encoding="utf-8"))[0] if source else {}
    tonality = bool(args.tonality or front.get("tonality"))
    info = docx_layout.read_docx(path)
    found = docx_layout.structural(info, front, args.kind == "manuscript", tonality) + docx_layout.restraint(path, tonality)
    result = {"components": info["components"], "rendered": False}
    if args.compare:
        mine, other = docx_layout.structure(path), docx_layout.structure(args.compare.expanduser().resolve())
        differ = docx_layout.structure_diff(mine, other)
        result["structure"] = {"this": mine, "other": other, "differ": differ}
        if len(differ) < 3:
            found.append(docx_layout.finding("tonality.structure", "FAIL", None,
                                             f"the two builds differ in {len(differ)} structural feature(s): {', '.join(differ) or 'none'}",
                                             "Make the tonalities differ in structure: title block, summary form, numbering, components, running head, contents."))
    if args.layout:
        with tempfile.TemporaryDirectory(prefix="lithermes-docx-pages-") as tmp:
            try:
                pdf = render_pdf(path, Path(tmp))
            except subprocess.TimeoutExpired:
                pdf = None
            if pdf is None:
                result["render"] = "LibreOffice could not render the document; the page checks did not run"
            else:
                page_found, pages = docx_layout.paged(info, docx_layout.read_pdf(pdf), tonality, front)
                found += page_found
                result.update(pages, rendered=True)
    result["findings"] = [f for f in found if f["severity"] == "FAIL"]
    result["advisories"] = [f for f in found if f["severity"] != "FAIL"]
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Pass/fail gate for a generated DOCX (LitHermes).")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--source", type=Path, help="the Markdown the document was built from")
    parser.add_argument("--publisher", help="korean-generic, elsevier, acs, ieee or nature; omit for a tonality or the plain build")
    parser.add_argument("--kind", choices=("report", "manuscript"), default="report")
    parser.add_argument("--tonality", help="the tonality used, when the source frontmatter does not name it")
    parser.add_argument("--layout", action="store_true", help="render with LibreOffice and check the pages")
    parser.add_argument("--compare", type=Path, help="the same source built in another tonality")
    args = parser.parse_args(argv)

    path = args.docx.expanduser().resolve()
    source = args.source.expanduser().resolve() if args.source else None
    problems = package_problems(path)
    ok = not problems
    content = content_report(path, source.read_text(encoding="utf-8") if source else None) if ok else {"pass": False}
    lint = lint_report(source, path, args.publisher or "korean-generic", args.kind, args.publisher is not None) if ok else {"pass": False}
    output = output_report(path, source, args) if ok else {"findings": [], "advisories": []}

    reasons = []
    if problems:
        reasons.append("package: " + "; ".join(problems[:5]))
    if content.get("placeholders"):
        reasons.append(f"content: unfilled blanks {content['placeholders'][:5]}; write realistic values and label them as examples")
    if content.get("frontmatter_leak"):
        reasons.append("content: frontmatter printed as body text")
    if ok and not content.get("pass") and not content.get("placeholders") and not content.get("frontmatter_leak"):
        reasons.append("content: empty document or headings missing from the output")
    if lint.get("prose"):
        reasons.append(f"prose lint: {len(lint['prose'])} finding(s)")
    if lint.get("design"):
        reasons.append(f"design audit: {len(lint['design'])} finding(s)")
    if output["findings"]:
        reasons.append("output checks: " + ", ".join(sorted({f["check"] for f in output["findings"]})))
    report = {"file": str(path), "pass": not reasons, "failure_reasons": reasons, "profile": args.publisher or "plain",
              "kind": args.kind, "package": problems, "content": content, "lint": lint, "output": output}
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
