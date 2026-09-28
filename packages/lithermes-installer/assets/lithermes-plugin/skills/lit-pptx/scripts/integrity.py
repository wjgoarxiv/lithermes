"""Small OOXML package integrity check owned by the LitHermes deck engine."""

from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from pptx import Presentation

PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT = "http://schemas.openxmlformats.org/package/2006/content-types"


def check_package(filename):
    path = Path(filename)
    with ZipFile(path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)) or archive.testzip() is not None:
            raise ValueError("damaged or duplicate OOXML archive member")
        members = set(names)
        if not {"[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml"} <= members:
            raise ValueError("PowerPoint package is missing core members")
        types = ET.fromstring(archive.read("[Content_Types].xml"))
        overrides = {part.attrib.get("PartName", "").lstrip("/") for part in types.findall(f"{{{CONTENT}}}Override")}
        if "ppt/presentation.xml" not in overrides:
            raise ValueError("presentation content type is missing")
        for member in names:
            part = PurePosixPath(member)
            if part.is_absolute() or ".." in part.parts:
                raise ValueError("unsafe OOXML member path")
            if not member.endswith(".rels"):
                continue
            root = ET.fromstring(archive.read(member))
            owner = PurePosixPath(member).parent.parent
            for rel in root.findall(f"{{{PKG}}}Relationship"):
                if rel.attrib.get("TargetMode") == "External":
                    continue
                target = rel.attrib.get("Target", "").split("#", 1)[0]
                if target.startswith("/"):
                    resolved = target.lstrip("/")
                else:
                    stack = list(owner.parts) if member != "_rels/.rels" else []
                    for segment in PurePosixPath(target).parts:
                        if segment == "..":
                            if not stack:
                                raise ValueError("relationship escapes package")
                            stack.pop()
                        elif segment != ".":
                            stack.append(segment)
                    resolved = "/".join(stack)
                if resolved not in members:
                    raise ValueError(f"broken OOXML relationship: {member} -> {resolved}")
    deck = Presentation(path)
    if not deck.slides:
        raise ValueError("PowerPoint package contains no slides")
    return {"slides": len(deck.slides), "members": len(names)}


if __name__ == "__main__":
    import json
    import sys
    print(json.dumps(check_package(sys.argv[1])))
