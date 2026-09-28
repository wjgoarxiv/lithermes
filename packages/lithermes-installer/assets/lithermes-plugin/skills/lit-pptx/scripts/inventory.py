"""Hermes deck layout inventory for geometry defects in text and pictures."""

import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def inspect(filename):
    deck = Presentation(filename)
    result = {}
    for number, slide in enumerate(deck.slides, 1):
        shapes = {}
        candidates = []
        for index, shape in enumerate(slide.shapes, 1):
            bounds = (shape.left, shape.top, shape.left + shape.width, shape.top + shape.height)
            kind = "picture" if shape.shape_type == MSO_SHAPE_TYPE.PICTURE else "text" if shape.has_text_frame else "shape"
            record = {"type": kind, "overflow": {}, "overlap": {}}
            if bounds[0] < -9144 or bounds[1] < -9144 or bounds[2] > deck.slide_width + 9144 or bounds[3] > deck.slide_height + 9144:
                record["overflow"]["slide"] = True
            if shape.has_text_frame and shape.text.strip():
                frame = shape.text_frame
                content_h = max(1, sum(max(1, p.text.count("\n") + 1) * (max((r.font.size.pt for r in p.runs if r.font.size), default=14) * 1.3) for p in frame.paragraphs))
                available_h = max(0, (shape.height - frame.margin_top - frame.margin_bottom) / 12700)
                if content_h > available_h * 1.35:
                    record["overflow"]["frame"] = True
            if kind in ("text", "picture"):
                candidates.append((index, bounds, record))
            shapes[f"shape-{index}"] = record
        for index, box, record in candidates:
            if record["type"] != "text":
                continue
            overlaps = []
            for other_index, other_box, other in candidates:
                if other_index == index or other["type"] != "picture":
                    continue
                area = max(0, min(box[2], other_box[2]) - max(box[0], other_box[0])) * max(0, min(box[3], other_box[3]) - max(box[1], other_box[1]))
                if area > 0.25 * max(1, (box[2] - box[0]) * (box[3] - box[1])):
                    overlaps.append(f"shape-{other_index}")
            if overlaps:
                record["overlap"]["overlapping_shapes"] = overlaps
        result[f"slide-{number}"] = shapes
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("deck")
    parser.add_argument("output", nargs="?")
    parser.add_argument("--issues-only", action="store_true")
    args = parser.parse_args()
    report = inspect(args.deck)
    if args.issues_only:
        report = {slide: {key: item for key, item in shapes.items() if item["overflow"] or item["overlap"]} for slide, shapes in report.items()}
    body = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(body, encoding="utf-8")
    else:
        print(body)


if __name__ == "__main__":
    main()
