"""Additive, non-destructive craft checks over the native slide object model."""

from __future__ import annotations

import colorsys
import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

from deck_output import review as review_deck_output


RULES = {
    "accent_saturation": .12, "accent_light_min": .12, "accent_light_max": .93,
    "hue_family_deg": 20, "full_bleed": .92, "body_min_pt": 10.5,
    "display_pt": 30, "latin_line_max": 90, "cjk_line_max": 38,
    "numeric_share": .70, "frame_inset_in": .06, "radius_error_in": .02,
    "outer_radius_min_in": .04, "group_ratio": 2,
    "empty_band_high": .35, "empty_band_medium": .25,
}
ADVISORY_BY_TEMPLATE = {"AZURE-PRO": {"OF-109"}}
EMU = 914400
NUMERIC = re.compile(r"^[\s,$€£¥₩%▲▼+\-±().0-9]+$")
EMOJI = re.compile("[\U0001f000-\U0001faff\u2600-\u27bf]")


def _box(shape):
    return tuple(float(getattr(shape, key)) / EMU for key in ("left", "top", "width", "height"))


def _rgb(shape):
    try:
        value = str(shape.fill.fore_color.rgb)
        return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))
    except (AttributeError, TypeError, ValueError):
        return None


def _hsl(rgb):
    h, l, s = colorsys.rgb_to_hls(*(value / 255 for value in rgb))
    return h * 360, s, l


def _radius(shape):
    if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or not shape.auto_shape_type:
        return None
    if "ROUND" not in str(shape.auto_shape_type).upper():
        return None
    guides = shape._element.xpath(".//a:avLst/a:gd[@name='adj']")
    if not guides:
        return None
    match = re.search(r"val\s+(\d+)", guides[0].get("fmla", ""))
    return float(match.group(1)) / 100000 * min(shape.width, shape.height) / EMU if match else None


def _bleeds(shape, width, height, tol=.03 * EMU):
    """A fill touching two canvas edges is a band, rail or caption field: page decoration, not a card."""
    left, top = shape.left, shape.top
    return sum([left <= tol, top <= tol, left + shape.width >= width - tol, top + shape.height >= height - tol]) >= 2


def _title_or_display(shapes):
    text = [shape for shape in shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
    if len(text) <= 1:
        return True
    return any(max((run.font.size.pt for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.size), default=0) >= RULES["display_pt"] for shape in text) and len(text) <= 2


def _numeric(value):
    value = value.strip()
    return bool(value and NUMERIC.fullmatch(value) and re.sub(r"[^0-9]", "", value))


def _contains(outer, inner):
    ox, oy, ow, oh = _box(outer)
    ix, iy, iw, ih = _box(inner)
    return inner is not outer and ox <= ix and oy <= iy and ix + iw <= ox + ow and iy + ih <= oy + oh


def _gap(first, second):
    ax, ay, aw, ah = _box(first)
    bx, by, bw, bh = _box(second)
    return ((max(0, ax - bx - bw, bx - ax - aw)) ** 2 +
            (max(0, ay - by - bh, by - ay - ah)) ** 2) ** .5


def check(path: Path) -> dict:
    deck = Presentation(str(path))
    subject = deck.core_properties.subject or ""
    template_id = subject.split("template:", 1)[1].strip() if "template:" in subject else ""
    findings = []

    def add(rule, severity, slide_no, shape, value, threshold, *, tier="measured"):
        advisory = rule in ADVISORY_BY_TEMPLATE.get(template_id, set())
        findings.append({"rule": rule, "severity": severity, "slide": slide_no,
                         "shape": getattr(shape, "name", ""), "value": value,
                         "threshold": threshold, "tier": tier,
                         **({"advisory_template": template_id} if advisory else {})})

    for slide_no, slide in enumerate(deck.slides, 1):
        shapes = list(slide.shapes)
        display = _title_or_display(shapes)
        # A pack slide names its family: a cover, section, statement, quote or full-bleed picture is a display slide.
        family = next((shape.name.split("@", 1)[1] for shape in shapes if shape.name.startswith("family@")), "")
        if family:
            display = bool(re.match(r"(cover|section)(-|$)|(statement|quote|closing-statement|image-full)$", family))
        page_area = deck.slide_width * deck.slide_height
        hues = []
        for shape in shapes:
            if shape.width * shape.height >= page_area * RULES["full_bleed"]:
                continue
            if re.search(r"circle-|template-image|deco", shape.name or "", re.I):
                continue
            color = _rgb(shape)
            if color is None:
                continue
            hue, saturation, lightness = _hsl(color)
            if saturation < RULES["accent_saturation"] or not (RULES["accent_light_min"] <= lightness <= RULES["accent_light_max"]):
                continue
            if not any(min(abs(hue - known), 360 - abs(hue - known)) <= RULES["hue_family_deg"] for known in hues):
                hues.append(hue)
        high = 4 if display else 3
        if len(hues) >= high:
            add("OF-101", "HIGH", slide_no, slide, len(hues), f"<{high} accent families")
        elif len(hues) == high - 1:
            add("OF-101", "MEDIUM", slide_no, slide, len(hues), f"<{high - 1} accent families")

        for shape in shapes:
            xml = shape._element
            if xml.xpath(".//a:rPr/a:gradFill"):
                add("OF-106", "HIGH", slide_no, shape, "gradient text", "none")
            if xml.xpath(".//a:effectLst/a:glow"):
                add("OF-107", "HIGH", slide_no, shape, "glow", "none")
            for bullet in xml.xpath(".//a:pPr/a:buChar"):
                if EMOJI.search(bullet.get("char", "")):
                    add("OF-108", "HIGH", slide_no, shape, bullet.get("char"), "non-emoji bullet")
            if getattr(shape, "has_table", False):
                table = shape.table
                if len(table.rows) >= 3:
                    for column in range(len(table.columns)):
                        cells = [table.cell(row, column) for row in range(1, len(table.rows))]
                        numeric = [cell for cell in cells if _numeric(cell.text)]
                        if len(numeric) / len(cells) < RULES["numeric_share"]:
                            continue
                        for cell in numeric:
                            for paragraph in cell.text_frame.paragraphs:
                                if paragraph.alignment is not None and paragraph.alignment != PP_ALIGN.RIGHT:
                                    add("OF-103", "HIGH", slide_no, shape, paragraph.alignment, "right-aligned numeric column")
                                elif paragraph.alignment is None:
                                    add("OF-103", "MEDIUM", slide_no, shape, "inherited alignment", "verified right alignment", tier="derived")
            # The body measure governs reading text; a display slide's lines are display type (pack decks only,
            # where the family says which slide is which).
            if getattr(shape, "has_text_frame", False) and shape.text.strip() and not (family and display):
                content = shape.text.strip()
                sizes = [run.font.size.pt for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.size]
                size = max(sizes, default=14)
                if RULES["body_min_pt"] < size < RULES["display_pt"] and shape.width > 0:
                    cjk = sum("가" <= char <= "힣" for char in content) / len(content) >= .5
                    capacity = max(1, (shape.width / EMU) * 72 / size)
                    # Each line the author or the engine broke wraps on its own; estimating one block hid the
                    # breaks and charged the break characters to the measure.
                    pieces = [piece for piece in re.split(r"[\n\v]", content) if piece.strip()]
                    lines = max(1, sum(max(1, round(sum(1 if "가" <= char <= "힣" else .55 for char in piece) / capacity))
                                       for piece in pieces))
                    printed = sum(len(piece) for piece in pieces)
                    ceiling = RULES["cjk_line_max"] if cjk else RULES["latin_line_max"]
                    if lines >= 2 and printed / lines > ceiling:
                        add("OF-102", "HIGH", slide_no, shape, round(printed / lines, 1), ceiling, tier="derived")

        rounded = [(shape, _radius(shape)) for shape in shapes]
        for outer, outer_radius in rounded:
            if outer_radius is None or outer_radius <= RULES["outer_radius_min_in"]:
                continue
            ox, oy, ow, oh = _box(outer)
            for inner, inner_radius in rounded:
                if inner is outer or inner_radius is None:
                    continue
                ix, iy, iw, ih = _box(inner)
                insets = (ix - ox, ox + ow - ix - iw, iy - oy, oy + oh - iy - ih)
                if min(insets) < 0 or max(insets) - min(insets) > RULES["frame_inset_in"]:
                    continue
                expected = max(0, outer_radius - sum(insets) / len(insets))
                if abs(inner_radius - expected) > RULES["radius_error_in"]:
                    add("OF-104", "HIGH", slide_no, inner, round(inner_radius, 3), round(expected, 3), tier="derived")

        cards = [shape for shape in shapes if _rgb(shape) is not None and
                 .05 * page_area <= shape.width * shape.height < RULES["full_bleed"] * page_area
                 and not _bleeds(shape, deck.slide_width, deck.slide_height)]
        groups = [(card, [child for child in shapes if _contains(card, child) and
                          (getattr(child, "has_text_frame", False) and child.text.strip() or getattr(child, "has_chart", False))])
                  for card in cards]
        groups = [(card, children) for card, children in groups if len(children) >= 2]
        if len(groups) >= 2:
            inner_gaps = [min(_gap(a, b) for i, a in enumerate(children) for b in children[i + 1:])
                          for _, children in groups]
            outer_gap = min(_gap(a, b) for i, (a, _) in enumerate(groups) for b, _ in groups[i + 1:])
            within = max(inner_gaps)
            if within > 0 and outer_gap < RULES["group_ratio"] * within:
                add("OF-105", "MEDIUM", slide_no, slide, round(outer_gap / within, 2), RULES["group_ratio"], tier="derived")
        trailing_by_card = []
        for card, children in ([] if family and display else groups):
            if _title_or_display(children):
                continue
            _, cy, _, ch = _box(card)
            last_bottom = max(_box(child)[1] + _box(child)[3] for child in children)
            trailing = max(0, cy + ch - last_bottom) / max(ch, .01)
            trailing_by_card.append((card, trailing))
        for card, trailing in trailing_by_card:
            _, cy, _, ch = _box(card)
            row = [value for peer, value in trailing_by_card
                   if abs(_box(peer)[1] - cy) <= .05 and abs(_box(peer)[3] - ch) <= .05]
            if len(row) > 1:
                trailing = min(row)
            severity = "HIGH" if trailing > RULES["empty_band_high"] else "MEDIUM" if trailing >= RULES["empty_band_medium"] else None
            if severity:
                add("OF-109", severity, slide_no, card, round(trailing, 3), RULES["empty_band_high"], tier="derived")

        # The content-area gap check runs only on genuine multi-block content slides.
        content = [shape for shape in shapes if (getattr(shape, "has_table", False) or getattr(shape, "has_chart", False)
                   or (getattr(shape, "has_text_frame", False) and len(shape.text.strip()) >= 20))
                   and shape.top > deck.slide_height * .16 and shape.width * shape.height < page_area * RULES["full_bleed"]]
        if not display and len(content) >= 2:
            bottoms = [shape.top + shape.height for shape in content]
            area_bottom = deck.slide_height * .88
            trailing = max(0, area_bottom - max(bottoms)) / max(1, deck.slide_height * .70)
            severity = "HIGH" if trailing > RULES["empty_band_high"] else "MEDIUM" if trailing >= RULES["empty_band_medium"] else None
            if severity:
                add("OF-109", severity, slide_no, slide, round(trailing, 3), RULES["empty_band_high"], tier="derived")

    # Deck-wide checks for pack-built decks: treatments, composition variety, body use, title frames, labels.
    findings.extend(review_deck_output(deck))
    return {"pass": not any(item["severity"] == "HIGH" and not item.get("advisory_template") for item in findings),
            "template_id": template_id, "findings": findings}
