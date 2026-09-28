"""Office layout contracts, exercised with the pinned Office runtime when present."""

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path


HAS_OFFICE = all(importlib.util.find_spec(name) for name in ("pptx", "docx", "markdown", "bs4", "yaml"))
SKILLS = Path(__file__).resolve().parents[2] / "assets/lithermes-plugin/skills"


@unittest.skipUnless(HAS_OFFICE, "run with the pinned Office Python runtime")
class OfficeQuality(unittest.TestCase):
    def load(self, path, name):
        spec = importlib.util.spec_from_file_location(name, path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    def test_sparse_table_slide_fails_craft_gate(self):
        from pptx import Presentation
        from pptx.util import Inches

        qa = self.load(SKILLS / "lit-pptx/scripts/qa_deck.py", "office_qa")
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(10), Inches(0.7)).text = "Quarterly comparison"
        table = slide.shapes.add_table(2, 3, Inches(0.8), Inches(2), Inches(11.7), Inches(0.8)).table
        table.cell(0, 0).text = "Metric"
        table.cell(1, 0).text = "Revenue"
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "sparse.pptx"
            prs.save(deck)
            craft = qa.check_slide_craft(deck)
        self.assertFalse(craft["pass"])
        self.assertTrue(any("table" in item["reason"] for item in craft["violations"]))

    def test_repeated_text_only_slides_fail_craft_gate(self):
        from pptx import Presentation
        from pptx.util import Inches

        qa = self.load(SKILLS / "lit-pptx/scripts/qa_deck.py", "office_qa_text")
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for index in range(4):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(0.7)).text = f"Section {index + 1}"
            slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11), Inches(4)).text = "The same text-only layout repeats across this presentation. It contains enough text to be content."
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "bullets.pptx"
            prs.save(deck)
            craft = qa.check_slide_craft(deck)
        self.assertFalse(craft["pass"])
        self.assertTrue(any("text-only" in item["reason"] for item in craft["violations"]))

    def test_docx_table_rows_do_not_split_and_columns_use_page_width(self):
        from docx import Document
        from docx.oxml.ns import qn

        converter = self.load(SKILLS / "lit-docx/scripts/convert_md_to_docx.py", "office_docx")
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        for row in table.rows:
            for cell in row.cells:
                cell.text = "A fairly long entry"
        converter.apply_table_style(doc, {})
        self.assertFalse(table.autofit)
        self.assertEqual(sum(col.width for col in table.columns), doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin)
        for row in table.rows:
            self.assertIsNotNone(row._tr.trPr.find(qn("w:cantSplit")))
            self.assertEqual([cell.width for cell in row.cells], [col.width for col in table.columns])


if __name__ == "__main__":
    unittest.main()
