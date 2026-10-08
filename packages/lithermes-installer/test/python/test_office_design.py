"""Design-direction checks for decks: the deck-wide output checks (OF-110..OF-114) and the engine's packs.

Run with the pinned Office runtime Python (the office cache venv); under the bare Hermes interpreter the
python-pptx imports are missing and the class is skipped.
"""

from __future__ import annotations

import importlib.util
import math
import os
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt
except ImportError as error:  # the Hermes interpreter has no Office runtime
    Presentation = None
    IMPORT_ERROR = str(error)
else:
    IMPORT_ERROR = ""

ROOT = Path(__file__).resolve().parents[2]
SKILL = Path(os.environ.get("LITHERMES_TEST_SKILLS", ROOT / "assets/lithermes-plugin/skills")) / "lit-pptx"
OFFICE = SKILL / "bin/office.mjs"
sys.path.insert(0, str(SKILL / "scripts"))


def _load(name):
    spec = importlib.util.spec_from_file_location(f"design_{name}", SKILL / "scripts" / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _runtime_ready():
    probe = subprocess.run(["node", str(OFFICE), "doctor"], capture_output=True, text=True)
    return "office-runtime=READY" in probe.stdout


def _text(slide, name, x, y, w, h, text, size):
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    box.name = name
    box.text_frame.text = text
    for run in box.text_frame.paragraphs[0].runs:
        run.font.size = Pt(size)
    return box


def _flat_deck(path, title_text="분기별 매출 추이"):
    """Eight slides that all wear one top-rule title over a short body pinned to the top."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    prs.core_properties.subject = "lit-pptx tonality=ledger density=10 grid=compact variance=5"
    for number in range(8):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 0, 0)
        marker.name = "family@content"
        _text(slide, "title@top-rule", 24, 36, 900, 40, title_text if number == 3 else f"분기 항목 {number + 1}", 26)
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(24), Pt(82), Pt(54), Pt(3))
        rule.fill.solid()
        rule.fill.fore_color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
        _text(slide, f"Text {number}", 24, 100, 600, 40, "짧은 본문 한 줄과 설명이 들어갑니다", 13)
    prs.save(path)


@unittest.skipIf(Presentation is None, f"python-pptx unavailable: {IMPORT_ERROR}")
class DeckOutputChecks(unittest.TestCase):
    def test_flat_tonality_deck_fails_treatment_variety_and_band_checks(self):
        craft = _load("craft_extras")
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "flat.pptx"
            _flat_deck(deck)
            result = craft.check(deck)
        rules = {item["rule"] for item in result["findings"] if item["severity"] == "HIGH"}
        self.assertIn("OF-110", rules)  # one title treatment on every slide
        self.assertIn("OF-111", rules)  # one composition everywhere
        self.assertIn("OF-112", rules)  # the lower body stays empty
        self.assertFalse(result["pass"])

    def test_declarative_title_is_a_label_finding(self):
        craft = _load("craft_extras")
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "claim.pptx"
            _flat_deck(deck, title_text="분기 매출이 처음으로 감소했다")
            labels = [item for item in craft.check(deck)["findings"] if item["rule"] == "OF-114"]
        self.assertEqual([item["slide"] for item in labels], [4])
        self.assertEqual(labels[0]["severity"], "HIGH")

    def test_label_rule_reads_korean_and_english_endings(self):
        output = _load("deck_output")
        for text in ("매출이 증가했다", "목표를 달성함", "Revenue grew in Q3", "Conversion is up."):
            self.assertTrue(output.is_claim(text), text)
        for text in ("분기별 매출 추이", "운영 자금 30억 원 요청 내역", "Conversion by feed rate, runs 1-12", "왜 지금인가?", "바다"):
            self.assertFalse(output.is_claim(text), text)

    def test_outlined_card_around_text_frames_is_not_an_empty_frame(self):
        qa = _load("qa_deck")
        with tempfile.TemporaryDirectory() as tmp:
            prs = Presentation()
            prs.slide_width, prs.slide_height = Pt(960), Pt(540)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            _text(slide, "title@top-rule", 24, 36, 900, 40, "품목별 대응 우선순위", 26)
            card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(414), Pt(63), Pt(249), Pt(117))
            card.line.color.rgb = RGBColor(0xC8, 0xD0, 0xDA)  # outline only, like a matrix quadrant
            _text(slide, "Text 8", 432, 81, 213, 23, "포장재 단가 재협상", 14)
            path = Path(tmp) / "card.pptx"
            prs.save(path)
            craft = qa.check_slide_craft(path)
        self.assertFalse([v for v in craft["violations"] if v["reason"] == "empty outlined shape"])

    def test_pack_slides_with_short_korean_takeaways_and_top_figures_are_not_flat_decks(self):
        qa = _load("qa_deck")
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as tmp:
            prs = Presentation()
            prs.slide_width, prs.slide_height = Pt(960), Pt(540)
            png = Path(tmp) / "fig.png"
            png.write_bytes(bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                                          "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082"))
            for number in range(4):
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 0, 0)
                if number < 3:
                    marker.name = "family@table-insight"
                    _text(slide, "title@top-rule", 24, 36, 900, 40, f"권역별 출고 실적 {number + 1}", 26)
                    slide.shapes.add_table(4, 3, Pt(24), Pt(100), Pt(560), Pt(200))
                    for row, line in enumerate(("수도권 출고가 석 달 연속 늘어 전체의 절반을 넘었다",
                                                "영남권은 물류센터 이전 뒤 처리 시간이 줄었다",
                                                "호남권은 반품 비중이 높아 별도 점검이 필요하다")):
                        _text(slide, f"Text {row}", 620, 100 + row * 40, 300, 30, line, 13)
                else:
                    marker.name = "family@figure-academic"
                    slide.shapes.add_picture(str(png), Pt(24), Pt(36), Pt(900), Pt(330))
                    _text(slide, "title@bottom-anchor", 24, 402, 640, 60, "시료별 흡착 곡선", 26)
            path = Path(tmp) / "pack.pptx"
            prs.save(path)
            craft = qa.check_slide_craft(path)
        reasons = {v["reason"] for v in craft["violations"]}
        self.assertNotIn("table-only deck", reasons)
        self.assertNotIn("text-only deck", reasons)

    def test_measure_reads_each_broken_line_and_still_flags_a_wide_frame(self):
        craft = _load("craft_extras")
        with tempfile.TemporaryDirectory() as tmp:
            prs = Presentation()
            prs.slide_width, prs.slide_height = Pt(960), Pt(540)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            _text(slide, "title@side-rail", 24, 36, 288, 90, "회귀 분석 절차", 26)
            _text(slide, "Text 1", 492, 67, 338, 55, "최소제곱법으로 직선을 구한다 · 도구 스프레드시트 회귀\n"
                  "함수 · 약 5분 · 결과 β₀ 51, β₁ 4.2 · 잔차 제곱합이\n가장 작은 직선", 13)
            wide = slide.shapes.add_textbox(Pt(24), Pt(300), Pt(900), Pt(60))
            wide.text_frame.word_wrap = True
            wide.text_frame.text = "가" * 70 + " " + "나" * 70
            wide.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
            path = Path(tmp) / "measure.pptx"
            prs.save(path)
            found = [f for f in craft.check(path)["findings"] if f["rule"] == "OF-102" and f["severity"] == "HIGH"]
        self.assertEqual([f["shape"] for f in found], ["TextBox 3"])

    def test_caption_band_bleeding_from_the_edge_is_not_a_hollow_card(self):
        craft = _load("craft_extras")
        with tempfile.TemporaryDirectory() as tmp:
            prs = Presentation()
            prs.slide_width, prs.slide_height = Pt(960), Pt(540)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Pt(960), Pt(540))
            marker.name = "family@image-full"
            band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Pt(324), Pt(570), Pt(192))
            band.fill.solid()
            band.fill.fore_color.rgb = RGBColor(0x12, 0x20, 0x33)
            _text(slide, "title@overlay", 24, 348, 522, 72, "4인 식탁으로 펼친 상태", 26)
            _text(slide, "Text 1", 24, 432, 522, 14, "도 1. 4인 식탁으로 펼친 모습, 상판 120 × 80cm | 출처: 예시 이미지", 11)
            path = Path(tmp) / "overlay.pptx"
            prs.save(path)
            found = [f for f in craft.check(path)["findings"] if f["rule"] == "OF-109"]
        self.assertEqual(found, [])

    def test_structured_pack_pages_are_not_bullet_pages_but_plain_text_pages_still_are(self):
        qa = _load("qa_deck")

        def deck(tmp, family):
            prs = Presentation()
            prs.slide_width, prs.slide_height = Pt(960), Pt(540)
            for number in range(4):
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 0, 0)
                marker.name = f"family@{family}"
                _text(slide, "title@top-rule", 24, 36, 900, 40, f"분기 운영 지표 {number + 1}", 26)
                for row in range(3):
                    _text(slide, f"Text {row}", 24 + row * 300, 120, 280, 60, "평균 처리 시간 3.4시간, 전 분기 3.9시간 대비 단축", 13)
            path = Path(tmp) / f"{family}.pptx"
            prs.save(path)
            return {v["reason"] for v in qa.check_slide_craft(path)["violations"]}

        with tempfile.TemporaryDirectory() as tmp:
            self.assertNotIn("text-only deck", deck(tmp, "kpi-row"))
            self.assertIn("text-only deck", deck(tmp, "text-column"))

    def test_title_named_by_treatment_counts_as_the_slide_title(self):
        lint = _load("validate_pptx")
        with tempfile.TemporaryDirectory() as tmp:
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            _text(slide, "title@side-rail", 24, 36, 288, 90, "원가 상승 품목별 대응", 26)
            path = Path(tmp) / "rail.pptx"
            prs.save(path)
            self.assertTrue(lint.lint(path)["checks"]["missing_titles"]["pass"])


@unittest.skipIf(Presentation is None, f"python-pptx unavailable: {IMPORT_ERROR}")
@unittest.skipUnless(_runtime_ready(), "office runtime cache not installed")
class DeckEngineDirections(unittest.TestCase):
    SOURCE = """---
title: 정비 부문 3분기 운영 점검
date: 2026-10-05
notice: 예시 데이터
---

---
layout: cover-typographic

# 정비 부문 3분기 운영 점검
---

---
layout: kpi-row

## 3분기 정비 핵심 지표

| 평균 수리 시간 | 재작업률 | 부품 대기 | 예방 정비 비중 |
|---|---|---|---|
| 3.4시간 | 2.1% | 1.8일 | 46% |
| 전 분기 3.9시간 | 전 분기 2.6% | 전 분기 2.4일 | 전 분기 41% |

- 수리 시간은 점검표 개편 뒤 줄었다
- 부품 대기는 공급사 두 곳 추가로 짧아졌다

출처: 예시 정비팀 작업 기록
---

---
layout: table-insight

## 설비군별 고장 건수

| 설비군 | 7월 | 8월 | 9월 |
|---|---:|---:|---:|
| 컨베이어 | 14 | 11 | 9 |
| 포장기 | 8 | 9 | 6 |
| 압축기 | 5 | 4 | 4 |
| 냉각기 | 3 | 5 | 2 |
| 지게차 | 6 | 6 | 5 |

- 컨베이어 고장이 석 달 연속 줄었다
- 냉각기는 8월 폭염 기간에만 늘었다

출처: 예시 설비 관리 시스템
---

---
layout: comparison

## 외주 정비와 내부 정비 비교

- 외주: 건당 비용 높음, 대기 2일
- 내부: 인력 추가 필요, 대기 0.5일
- 외주: 야간 대응 불가
- 내부: 야간 교대 가능
---

---
layout: closing-ask

## 4분기 요청 사항

| 항목 | 금액 | 시점 |
|---|---:|---|
| 예비 부품 확보 | 1.2억 원 | 10월 |
| 점검 인력 1명 | 0.6억 원 | 11월 |
---
"""

    def _compile(self, tmp, *flags):
        source = Path(tmp) / "deck.md"
        source.write_text(self.SOURCE, encoding="utf-8")
        out = Path(tmp) / f"deck{len(flags)}{abs(hash(flags))}.pptx"
        run = subprocess.run(["node", str(OFFICE), "pptx", str(source), *flags, "--pptx", str(out)],
                             capture_output=True, text=True, env={k: v for k, v in os.environ.items() if k != "TYPESAFE_API_KEY"})
        self.assertEqual(run.returncode, 0, run.stderr[-2000:])
        return out

    def test_unknown_tonality_is_a_clear_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "deck.md"
            source.write_text(self.SOURCE, encoding="utf-8")
            run = subprocess.run(["node", str(OFFICE), "pptx", str(source), "--tonality", "nonesuch", "--pptx", str(Path(tmp) / "x.pptx")],
                                 capture_output=True, text=True)
        self.assertNotEqual(run.returncode, 0)
        self.assertRegex(run.stderr + run.stdout, re.compile(r"nonesuch[\s\S]*ledger", re.I))

    def test_two_tonalities_draw_different_titles_and_stamp_their_pack(self):
        with tempfile.TemporaryDirectory() as tmp:
            ledger = Presentation(str(self._compile(tmp, "--tonality", "ledger")))
            signal = Presentation(str(self._compile(tmp, "--tonality", "signal")))
        def treatments(prs):
            return {shape.name for slide in prs.slides for shape in slide.shapes if shape.name.startswith("title@")}
        self.assertIn("tonality=ledger", ledger.core_properties.subject)
        self.assertIn("tonality=signal", signal.core_properties.subject)
        self.assertNotEqual(treatments(ledger), treatments(signal))

    def test_legacy_template_still_compiles_to_the_same_slide_count(self):
        with tempfile.TemporaryDirectory() as tmp:
            legacy = Presentation(str(self._compile(tmp, "--template", "AZURE-PRO")))
        self.assertEqual(len(legacy.slides), 5)
        self.assertIn("LitHermes template:AZURE-PRO", legacy.core_properties.subject)

    def test_tonality_deck_passes_the_output_checks(self):
        craft = _load("craft_extras")
        with tempfile.TemporaryDirectory() as tmp:
            result = craft.check(self._compile(tmp, "--tonality", "ledger"))
        high = [item for item in result["findings"] if item["severity"] == "HIGH" and item["rule"] in {"OF-110", "OF-113", "OF-114"}]
        self.assertEqual(high, [])

    def test_contact_split_closing_with_a_total_keeps_the_next_step_off_it(self):
        source = """---
title: 정비 거점 확충 요청
tonality: studio
notice: 예시 데이터
---

---
layout: cover-typographic

# 정비 거점 확충 요청
---

---
layout: closing-contact-split

## 2027년 정비 거점 확충 예산

| 쓰임 | 금액 |
|---|---|
| 거점 임차 | 6억 원 |
| 정비 장비 | 8억 원 |
| 정비 인력 | 5억 원 |
| 부품 재고 | 3억 원 |

- 다음 단계: 12월 셋째 주 예산 심의에서 거점별 개소 순서와 분기별 집행 한도를 함께 확정, 담당 예시 시설관리팀
---
"""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "closing.md"
            path.write_text(source, encoding="utf-8")
            out = Path(tmp) / "closing.pptx"
            run = subprocess.run(["node", str(OFFICE), "pptx", str(path), "--pptx", str(out)],
                                 capture_output=True, text=True, env={k: v for k, v in os.environ.items() if k != "TYPESAFE_API_KEY"})
            self.assertEqual(run.returncode, 0, run.stderr[-2000:])
            slide = Presentation(str(out)).slides[1]
        boxes = {}
        for shape in slide.shapes:
            text = shape.text_frame.text if shape.has_text_frame else ""
            for key in ("합계", "다음 단계"):
                if text.startswith(key):
                    boxes[key] = [round(v / 12700) for v in (shape.left, shape.top, shape.width, shape.height)]
        self.assertEqual(sorted(boxes), ["다음 단계", "합계"], boxes)
        (sx, sy, sw, sh), (tx, ty, tw, th) = boxes["다음 단계"], boxes["합계"]
        apart = sx + sw <= tx or tx + tw <= sx or sy + sh <= ty or ty + th <= sy
        self.assertTrue(apart, f"next step {boxes['다음 단계']} overlaps total {boxes['합계']}")

    def test_multi_line_chart_drops_point_values_and_horizontal_bars_keep_level_categories(self):
        source = """---
title: 정비 대기 시간 점검
tonality: night
notice: 예시 데이터
---

---
layout: cover-typographic

# 정비 대기 시간 점검
---

---
layout: content

## 공장별 월간 정비 대기 시간

::: chart type=line unit="시간"
| 월 | 1공장 | 2공장 | 3공장 |
|---|---|---|---|
| 7월 | 5.2 | 6.1 | 7.4 |
| 8월 | 4.9 | 5.8 | 7.0 |
| 9월 | 4.6 | 5.9 | 6.6 |
> 공장별 정비 대기 시간 (예시 데이터)
:::
출처: 예시 정비팀 작업 기록
---

---
layout: content

## 원인별 설비 정지 건수

::: chart type=bar unit="건"
| 원인 | 건수 |
|---|---|
| 윤활유 교체 주기 초과 | 7 |
| 센서 오염으로 인한 오작동 | 5 |
| 교대 인수인계 누락 | 4 |
| 부품 대기 | 2 |
> 원인별 설비 정지 건수 (예시 데이터)
:::
출처: 예시 설비 관리 시스템
---
"""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "charts.md"
            path.write_text(source, encoding="utf-8")
            out = Path(tmp) / "charts.pptx"
            run = subprocess.run(["node", str(OFFICE), "pptx", str(path), "--pptx", str(out)],
                                 capture_output=True, text=True, env={k: v for k, v in os.environ.items() if k != "TYPESAFE_API_KEY"})
            self.assertEqual(run.returncode, 0, run.stderr[-2000:])
            with zipfile.ZipFile(out) as package:
                charts = [package.read(n).decode() for n in package.namelist() if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)]
        line = next(x for x in charts if "<c:lineChart>" in x)
        bar = next(x for x in charts if '<c:barDir val="bar"/>' in x)
        # Three lines a few hours apart: a value printed over every point would sit on the neighbouring lines' values.
        self.assertNotIn('<c:showVal val="1"/>', line, "values of three line series print over each other")
        self.assertRegex(line, re.compile(r"<c:valAx>[\s\S]*?<c:delete val=\"0\"/>"), "without point values the value axis must show")
        category_axis = bar[bar.index("<c:catAx>"):bar.index("</c:catAx>")]
        self.assertNotIn('rot="-2700000"', category_axis, "horizontal bar categories are slanted")


    def _charts_of(self, body, tonality="night"):
        source = f"---\ntitle: 정비 지표 점검\ntonality: {tonality}\nnotice: 예시 데이터\n---\n\n---\n{body}\n---\n"
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "charts.md"
            path.write_text(source, encoding="utf-8")
            out = Path(tmp) / "charts.pptx"
            run = subprocess.run(["node", str(OFFICE), "pptx", str(path), "--pptx", str(out)],
                                 capture_output=True, text=True, env={k: v for k, v in os.environ.items() if k != "TYPESAFE_API_KEY"})
            self.assertEqual(run.returncode, 0, run.stderr[-2000:])
            with zipfile.ZipFile(out) as package:
                return [package.read(n).decode() for n in sorted(package.namelist()) if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)]

    def test_negative_horizontal_bars_put_category_labels_at_the_axis_minimum(self):
        (chart,) = self._charts_of("""layout: content

## 설비군별 정비 비용 증감

::: chart type=bar unit="천만 원"
| 설비군 | 증감 (천만 원) |
|---|---|
| 컨베이어 | 9 |
| 포장기 | 4 |
| 압축기 | -3 |
| 냉각기 | -6 |
> 설비군별 정비 비용 증감 (예시 데이터)
:::
출처: 예시 정비 회계 자료""", tonality="ledger")
        category_axis = chart[chart.index("<c:catAx>"):chart.index("</c:catAx>")]
        # Labels beside the zero line would sit on the two negative bars.
        self.assertIn('<c:tickLblPos val="low"/>', category_axis)

    def test_missing_cells_leave_gaps_in_a_line_instead_of_zeros(self):
        (chart,) = self._charts_of("""layout: content

## 공장별 월간 설비 가동률

::: chart type=line unit="%"
| 월 | 1공장 (%) | 2공장 (%) |
|---|---|---|
| 6월 | 91.2 | 88.4 |
| 7월 | 92.0 | 89.1 |
| 8월 | 92.6 | — |
| 9월 | 93.1 |  |
> 공장별 설비 가동률 (예시 데이터)
:::
출처: 예시 설비 관리 시스템""")
        second = chart[chart.rindex("<c:ser>"):]
        values = re.findall(r'<c:pt idx="(\d)"><c:v>([^<]*)</c:v>', second[second.index("<c:val>"):])
        drawn = [(idx, value) for idx, value in values if int(idx) >= 2 and value != ""]
        self.assertEqual(drawn, [], f"plant 2 has no August or September figure, yet points are drawn: {values}")
        self.assertIn('<c:dispBlanksAs val="gap"/>', chart)

    def test_thin_clustered_bars_drop_their_labels_and_a_narrow_bar_chart_keeps_few_axis_numbers(self):
        (pair,) = self._charts_of("""layout: kpi-over-chart

## 3분기 정비 요청 처리

| 처리 건수 | 예방 정비 비중 | 평균 처리 기간 |
|---|---|---|
| 864건 | 46% | 1.9일 |
| 2분기 790건 | 2분기 41% | 2분기 2.4일 |

::: chart type=bar unit="건"
| 요청 경로 | 2분기 | 3분기 |
|---|---|---|
| 현장 신고 | 310 | 280 |
| 센서 경보 | 190 | 330 |
| 정기 점검 | 160 | 170 |
| 협력사 요청 | 80 | 50 |
| 기타 | 50 | 34 |
> 요청 경로별 처리 건수 (예시 데이터)
:::

- 센서 경보가 처음으로 현장 신고를 넘었다
출처: 예시 정비팀 작업 기록""")
        self.assertNotIn('<c:showVal val="1"/>', pair, "value labels of two thin bars per row touch")
        self.assertRegex(pair, re.compile(r"<c:valAx>[\s\S]*?<c:delete val=\"0\"/>"))
        grid = self._charts_of("""layout: dashboard-grid

## Repair flow, uptime and response

::: chart type=bar unit="%"
| Stage | Share of requests (%) |
|---|---|
| Logged | 100 |
| Parts reserved | 74 |
| Technician assigned | 61 |
| Closed same week | 44 |
> Repair request stages (sample data)
:::

::: chart type=line unit="%"
| Month | Uptime (%) |
|---|---|
| Jul | 97.8 |
| Aug | 98.1 |
| Sep | 98.4 |
> Line uptime (sample data)
:::

::: chart type=column unit="h"
| Month | Response (h) |
|---|---|
| Jul | 3.9 |
| Aug | 3.6 |
| Sep | 3.4 |
> Median response time (sample data)
:::

- Parts reservation loses a quarter of requests""")
        stages = next(x for x in grid if '<c:barDir val="bar"/>' in x)
        unit = re.search(r'<c:valAx>[\s\S]*?<c:majorUnit val="([\d.]+)"/>', stages)
        self.assertIsNotNone(unit, "a narrow horizontal bar chart leaves its axis steps to LibreOffice, which crowds or slants them")
        # Long stage names leave the axis a narrow strip: four numbers fit there level.
        self.assertLessEqual(math.ceil(100 / float(unit.group(1))) + 1, 4)


if __name__ == "__main__":
    unittest.main()
