"""Direct tests for the places where the prose lint and the deck validator read a pattern as data, not as a defect.

Each case pairs the allowed form with the defect it must still catch. Run with the pinned Office runtime
Python (the office cache venv); under the bare Hermes interpreter the python-docx and python-pptx imports
are missing and the classes are skipped.
"""

import importlib.util
import random
import sys
import tempfile
import unittest
from pathlib import Path


HAS_OFFICE = all(importlib.util.find_spec(name) for name in ("pptx", "docx", "markdown", "bs4", "yaml"))
SKILLS = Path(__file__).resolve().parents[2] / "assets/lithermes-plugin/skills"


def _load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module  # a dataclass looks its module up by name
    spec.loader.exec_module(module)
    return module


@unittest.skipUnless(HAS_OFFICE, "run with the pinned Office Python runtime")
class ProseLintReadsDataAsData(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        scripts = SKILLS / "lit-docx/scripts"
        sys.path.insert(0, str(scripts))
        cls.lint = _load(scripts / "slop_lint.py", "relaxation_slop_lint")
        cls.rules = cls.lint.load_phrase_rules(scripts.parent / "references/slop_phrase_list.yaml")
        cls.publisher = cls.lint.load_publisher(scripts.parent / "templates/registry.yaml", "korean-generic")

    def rule_ids(self, text):
        return {finding.rule_id for finding in self.lint.lint_text(text, self.publisher, self.rules, "auto")[0]}

    def test_table_rows_are_exempt_from_the_dash_rule(self):
        table = ("# 목표\n\n성과 지표는 다음 표와 같습니다.\n\n| 지표 | 도입 전 | 목표 |\n| --- | ---: | ---: |\n"
                 "| 절감액 | — | 3.2 |\n| 가동률 | — | 95.0 |\n| 대기 시간 | 4 - 6 | 2 |\n")
        self.assertNotIn("rule-02-em-dash-cluster", self.rule_ids(table))
        prose = "# 목표\n\n절감액은 크다 — 그러나 가동률은 낮다 — 그래서 대기가 길다.\n"
        self.assertIn("rule-02-em-dash-cluster", self.rule_ids(prose))

    def test_a_hyphen_opening_a_line_is_a_list_marker_not_a_dash(self):
        listed = "# 핵심 수치\n\n처리 시간이 줄었습니다.\n\n- 전년 동기 4.1시간\n  - 목표 3.0시간\n- 재작업률 2.4%\n"
        self.assertNotIn("rule-02-em-dash-cluster", self.rule_ids(listed))
        spaced = "# 핵심 수치\n\n처리 시간이 줄었습니다 - 재작업률도 함께 내려갔습니다.\n"
        self.assertIn("rule-02-em-dash-cluster", self.rule_ids(spaced))

    def test_lexical_diversity_is_read_per_subsection(self):
        rng = random.Random(7)
        vocab = [f"term{i}" for i in range(80)]
        varied = "# Results\n\n" + "\n\n".join(f"## Part {k}\n\n" + " ".join(rng.sample(vocab, 60)) + "." for k in range(5)) + "\n"
        whole = varied.split("\n", 2)[2]
        self.assertLess(self.lint.type_token_ratio(whole), 0.45, "read as one block the section would be flagged")
        self.assertNotIn("rule-10-lexical-diversity", self.rule_ids(varied))
        flat = "# Results\n\n## Part one\n\n" + " ".join(["the loop and the catalyst and the loop"] * 12) + ".\n"
        self.assertIn("rule-10-lexical-diversity", self.rule_ids(flat))

    def test_lexical_diversity_is_skipped_when_no_subsection_reaches_forty_words(self):
        short = "# Results\n\n" + "\n\n".join(f"## Part {k}\n\n" + " ".join(["the loop and the loop"] * 6) + "." for k in range(3)) + "\n"
        parts = short.split("## ")[1:]
        self.assertTrue(all(self.lint.word_count("## " + part) <= 40 for part in parts))
        self.assertGreater(self.lint.word_count(short), 40)
        self.assertNotIn("rule-10-lexical-diversity", self.rule_ids(short))


@unittest.skipUnless(HAS_OFFICE, "run with the pinned Office Python runtime")
class DeckValidatorRelaxations(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.validate = _load(SKILLS / "lit-pptx/scripts/validate_pptx.py", "relaxation_validate_pptx")

    def test_an_acronym_ban_matches_the_whole_uppercase_word_only(self):
        pattern = self.validate._build_regex(["RAM", "synergy"], case_insensitive=True, whole_word=False)
        for text in ("RAM 사용량이 늘었습니다", "16 GB of RAM.", "(RAM)"):
            self.assertIsNotNone(pattern.search(text), text)
        for text in ("diagram", "program review", "RAMP", "ram usage", "Ramadan"):
            self.assertIsNone(pattern.search(text), text)
        self.assertIsNotNone(pattern.search("Real Synergy"), "a term that is not an acronym keeps the file's case setting")

    def test_a_closing_references_slide_counts_as_the_source_line(self):
        from pptx import Presentation
        from pptx.util import Pt

        def last_slide(*frames):
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(Pt(24), Pt(24), Pt(600), Pt(60)).text = "정비 운영 점검"
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            for row, text in enumerate(frames):
                slide.shapes.add_textbox(Pt(24), Pt(24 + row * 80), Pt(600), Pt(60)).text = text
            with tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "deck.pptx"
                prs.save(path)
                return self.validate.lint(path)["checks"]["missing_source_note"]["pass"]

        entries = "[1] Example Author A. Placeholder title. Example Journal. 2021;14:201-215."
        self.assertTrue(last_slide("References", entries))
        self.assertTrue(last_slide("참고 문헌", entries))
        self.assertTrue(last_slide("다음 단계", "출처: 정비 이력 시스템"))
        self.assertFalse(last_slide("References are listed in the appendix", "다음 분기에 다시 점검합니다"))
        self.assertFalse(last_slide("다음 단계", "다음 분기에 다시 점검합니다"))


if __name__ == "__main__":
    unittest.main()
