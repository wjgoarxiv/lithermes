"""The deck validator's word list ships with the slide skill, and a deck that keeps a placeholder is flagged.

validate_pptx.py reads lit-pptx/FORBIDDEN_TERMS.json; without the file its forbidden-term check finds nothing.
The list test runs under any interpreter; the deck test needs the pinned Office runtime (python-pptx).
"""

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path


SKILL = Path(__file__).resolve().parents[2] / "assets/lithermes-plugin/skills/lit-pptx"
try:
    from pptx import Presentation
    from pptx.util import Pt
except ImportError as error:  # an interpreter without the Office runtime (or with half of it)
    Presentation = None
    IMPORT_ERROR = str(error)
else:
    IMPORT_ERROR = ""


class ForbiddenTermList(unittest.TestCase):
    def test_the_word_list_ships_with_hard_and_soft_terms(self):
        path = SKILL / "FORBIDDEN_TERMS.json"
        self.assertTrue(path.is_file(), "lit-pptx/FORBIDDEN_TERMS.json is missing, so the forbidden-term check finds nothing")
        data = json.loads(path.read_text(encoding="utf-8"))
        self.assertTrue(data.get("terms"), "the hard term list is empty")
        self.assertTrue(data.get("soft_terms"), "the soft term list is empty")
        for key in ("terms", "soft_terms"):
            self.assertTrue(all(isinstance(term, str) and term.strip() for term in data[key]), f"{key} holds an empty entry")
        self.assertIn("lorem ipsum", [term.lower() for term in data["terms"]])


@unittest.skipIf(Presentation is None, f"python-pptx unavailable: {IMPORT_ERROR}")
class PlaceholderDeckIsFlagged(unittest.TestCase):
    def lint(self, *texts):
        spec = importlib.util.spec_from_file_location("forbidden_validate_pptx", SKILL / "scripts/validate_pptx.py")
        module = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = module
        spec.loader.exec_module(module)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for row, text in enumerate(texts):
            slide.shapes.add_textbox(Pt(24), Pt(24 + row * 80), Pt(600), Pt(60)).text = text
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            prs.save(path)
            return module.lint(path)

    def test_a_placeholder_term_fails_the_deck(self):
        report = self.lint("분기 운영 점검", "Lorem ipsum dolor sit amet")
        self.assertFalse(report["checks"]["forbidden_terms"]["pass"])
        self.assertFalse(report["pass"])
        self.assertTrue(any("Lorem ipsum" in hit for hit in report["checks"]["forbidden_terms"]["hard_hits"]))

    def test_a_single_word_term_inside_a_longer_word_passes(self):
        for text in ("Mastodon 서버 이전 계획", "photodocument scanner rollout"):
            report = self.lint("분기 운영 점검", text)
            self.assertTrue(report["checks"]["forbidden_terms"]["pass"], f"{text!r}: {report['checks']['forbidden_terms']['hard_hits']}")

    def test_a_single_word_term_standing_as_a_word_still_fails(self):
        for text in ("TODO: x", "todo.", "(TBD)", "Lorem ipsum", "FIXME 확인", "placeholder", "TODOs", "placeholders", "FIXMEs 정리"):
            report = self.lint("분기 운영 점검", text)
            self.assertFalse(report["checks"]["forbidden_terms"]["pass"], text)

    def test_plain_content_passes_the_term_check(self):
        report = self.lint("분기 운영 점검", "평균 수리 시간 3.4시간, 전 분기 3.9시간 대비 단축")
        self.assertTrue(report["checks"]["forbidden_terms"]["pass"])


if __name__ == "__main__":
    unittest.main()
