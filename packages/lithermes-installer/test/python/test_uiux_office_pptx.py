"""Regression fixtures for the additive slide craft checks."""

from __future__ import annotations

import importlib.util
import tempfile
import unittest
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except (ImportError, AttributeError) as error:
    Presentation = None
    OPTIONAL_IMPORT_ERROR = str(error)
else:
    OPTIONAL_IMPORT_ERROR = ""


ROOT = Path(__file__).resolve().parents[2]
QA = ROOT / "assets/lithermes-plugin/skills/lit-pptx/scripts/qa_deck.py"
sys.path.insert(0, str(QA.parent))
spec = importlib.util.spec_from_file_location("office_qa_deck", QA)
qa_deck = importlib.util.module_from_spec(spec) if Presentation is not None else None
if qa_deck is not None:
    spec.loader.exec_module(qa_deck)


@unittest.skipIf(Presentation is None, f"python-pptx unavailable: {OPTIONAL_IMPORT_ERROR}")
class OfficePptxChecks(unittest.TestCase):
    def test_accent_and_numeric_alignment_are_flagged_without_mutating_deck(self):
        with tempfile.TemporaryDirectory() as directory:
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for index, color in enumerate(((200, 40, 40), (20, 150, 40), (30, 70, 210))):
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1 + index), Inches(1), Inches(.8), Inches(.8))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor(*color)
            table = slide.shapes.add_table(3, 2, Inches(1), Inches(3), Inches(5), Inches(2)).table
            for row, text in enumerate(("Value", "123", "456")):
                cell = table.cell(row, 1)
                cell.text = text
                cell.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
            path = Path(directory) / "sample.pptx"
            deck.save(path)
            before = path.read_bytes()
            result = qa_deck.check_office_craft(path)
            ids = {item["rule"] for item in result["findings"]}
            self.assertIn("OF-101", ids)
            self.assertIn("OF-103", ids)
            self.assertFalse(result["pass"])
            self.assertEqual(path.read_bytes(), before)

    def test_plain_slide_has_no_high_new_craft_findings(self):
        with tempfile.TemporaryDirectory() as directory:
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
            text.text = "A clear, short statement"
            text.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
            path = Path(directory) / "plain.pptx"
            deck.save(path)
            result = qa_deck.check_office_craft(path)
            self.assertFalse(any(item["severity"] == "HIGH" for item in result["findings"]))

    def test_template_ornament_issue_is_advisory_for_recorded_template_id(self):
        with tempfile.TemporaryDirectory() as directory:
            deck = Presentation()
            deck.core_properties.subject = "LitHermes template:AZURE-PRO"
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for index in range(3):
                shape = slide.shapes.add_textbox(Inches(1), Inches(1.7 + index * .5), Inches(5), Inches(.4))
                shape.text = "A detailed content statement in the upper section"
            path = Path(directory) / "template.pptx"
            deck.save(path)
            result = qa_deck.check_office_craft(path)
            gap = [item for item in result["findings"] if item["rule"] == "OF-109"]
            self.assertTrue(gap)
            self.assertTrue(all(item["advisory_template"] == "AZURE-PRO" for item in gap))
            self.assertTrue(result["pass"])


if __name__ == "__main__":
    unittest.main()
