"""Document directions: the six tonality packs, page components, restraint rules and the document gate.

Run with the pinned Office runtime Python; under the bare Hermes interpreter python-docx is missing and the
classes are skipped.
"""

from __future__ import annotations

import importlib.util
import json
import re
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

try:
    import docx  # noqa: F401
    import yaml  # noqa: F401
except ImportError as error:
    HAVE_RUNTIME = False
    IMPORT_ERROR = str(error)
else:
    HAVE_RUNTIME = True
    IMPORT_ERROR = ""

ROOT = Path(__file__).resolve().parents[2]
SKILL = Path(os.environ.get("LITHERMES_TEST_SKILLS", ROOT / "assets/lithermes-plugin/skills")) / "lit-docx"
SCRIPTS = SKILL / "scripts"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

SOURCE = """---
title: 공용 자전거 정비 거점 1년 운영 결과
subtitle: 시범 구역 세 곳의 이용과 비용
short_title: 정비 거점 운영 결과
author: 예시 시설관리팀
organization: 예시 도시공사
date: 2026-10-01
notice: "예시 데이터: 기관과 수치는 모두 가정입니다"
---

# 요약

정비 거점 세 곳을 1년 동안 운영한 결과 고장 신고 처리 시간이 줄었다.

::: keyfigures
- **−38%** 평균 처리 시간
  - 2026년 9월, 도입 전 4.1일 대비
- **1,240건** 연간 정비 건수
  - 2025년 10월 – 2026년 9월
- **2.3억 원** 연간 운영비
  - 인건비와 부품비 합계
:::

::: callout kind=key title="결정 요청"
거점 두 곳을 추가할지 2026. 11. 30.까지 결정이 필요하다.
:::

# 운영 현황

거점별 정비 건수와 처리 시간은 다음과 같다.

표 1. 거점별 정비 실적 (예시)

| 거점 | 정비 건수 | 평균 처리 시간 | 부품비 |
|---|---:|---:|---:|
| 북부 | 520건 | 2.4일 | 3,100만 원 |
| 중앙 | 430건 | 2.6일 | 2,700만 원 |
| 남부 | 290건 | 2.9일 | 1,900만 원 |

자료: 예시 도시공사 정비 기록

## 고장 유형

체인과 브레이크 고장이 전체의 61%를 차지했다. 계절에 따라 타이어 고장이 늘었다.

::: sidebar title="용어"
- 처리 시간: 신고부터 재배치까지
- 거점: 정비 인력이 상주하는 시설
:::

# 비용과 효과

연간 운영비는 2.3억 원이고 외주 정비 대비 0.4억 원이 적었다.

## 향후 과제

부품 재고 관리와 야간 신고 대응이 남은 과제다.
"""


def _convert(tmp, *flags, source=SOURCE, name=None):
    src = Path(tmp) / "source.md"
    src.write_text(source, encoding="utf-8")
    out = Path(tmp) / f"{name or '-'.join(f.strip('-') for f in flags) or 'plain'}.docx"
    run = subprocess.run([sys.executable, str(SCRIPTS / "convert_md_to_docx.py"), str(src), str(out), *flags],
                         capture_output=True, text=True, env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
    return run, out, src


def _xml(path, part="word/document.xml"):
    with zipfile.ZipFile(path) as archive:
        return archive.read(part).decode("utf-8")


def _load(name):
    sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(f"docx_{name}", SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module  # dataclasses resolve their module through sys.modules
    spec.loader.exec_module(module)
    return module


@unittest.skipUnless(HAVE_RUNTIME, f"office runtime unavailable: {IMPORT_ERROR}")
class DocumentDirections(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="lithermes-docx-test-")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def build(self, *flags, **kw):
        run, out, src = _convert(self.tmp, *flags, **kw)
        self.assertEqual(run.returncode, 0, run.stderr[-1500:] + run.stdout[-500:])
        return out, src, run.stdout

    def test_six_packs_load_with_restrained_tokens_and_unknown_names_fail(self):
        design = _load("docx_design")
        names = sorted(design.TONALITIES)
        self.assertEqual([n.lower() for n in names], ["brief", "journal", "manual", "memo", "proposal", "report"])
        for name in names:
            pack = design.load_tonality(name)
            self.assertTrue(10.5 <= pack.body_pt <= 11, name)
            self.assertLessEqual(pack.sizes["h1"] / pack.body_pt, 1.5, name)
            self.assertLessEqual(len(pack.design.get("accent_on") or []), 2, name)
            self.assertNotIn("pullquote", pack.allowed_kinds, name)
            self.assertLessEqual(len(pack.allowed_kinds), 3, name)
            top, bottom, left, right = pack.margins_mm
            self.assertTrue(left >= 25 and right >= 25 and bottom >= top, name)
            self.assertAlmostEqual(pack.leading, pack.pitch / 1.53, places=2)
        self.assertEqual(design.load_tonality("memo").allowed_kinds, [])
        with self.assertRaises(design.DesignError):
            design.load_tonality("Glossy")
        with self.assertRaises(design.DesignError):
            design.load_tonality("report", density=11)

    def test_directives_become_components_and_never_print_as_text(self):
        out, _src, stdout = self.build("--tonality", "Report")
        printed = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", _xml(out)))
        for leak in (":::", "keyfigures", "column-break", "{style="):
            self.assertNotIn(leak, printed, leak)
        layout = _load("docx_layout")
        info = layout.read_docx(out)
        self.assertGreaterEqual(info["components"].get("keyfigures", 0), 1)
        self.assertGreaterEqual(info["components"].get("callout", 0), 1)
        self.assertEqual(info["components"].get("sidebar", 0), 0, "a report sets no sidebar")
        manual, _s, _o = self.build("--tonality", "Manual", name="manual")
        self.assertGreaterEqual(layout.read_docx(manual)["components"].get("sidebar", 0), 1)

    def test_tonality_pairs_differ_in_structure_not_only_colour(self):
        layout = _load("docx_layout")
        built = {name: self.build("--tonality", name, name=name)[0] for name in ("Report", "Brief", "Proposal", "Memo", "Manual")}
        for a in built:
            for b in built:
                if a < b:
                    differ = layout.structure_diff(layout.structure(built[a]), layout.structure(built[b]))
                    self.assertGreaterEqual(len(differ), 3, f"{a} vs {b}: {differ}")

    def test_restraint_holds_on_every_tonality(self):
        layout = _load("docx_layout")
        for name in ("Report", "Brief", "Proposal", "Memo", "Manual"):
            out = self.build("--tonality", name, name=f"r-{name}")[0]
            fails = [f for f in layout.restraint(out, True) if f["severity"] == "FAIL"]
            self.assertEqual(fails, [], name)

    def test_korean_document_declares_pretendard_a4_word_wrap_and_korean_conventions(self):
        out = self.build("--tonality", "Report")[0]
        styles, body = _xml(out, "word/styles.xml"), _xml(out)
        self.assertIn('w:ascii="Pretendard"', styles)
        self.assertIn('w:eastAsia="ko-KR"', styles)
        self.assertIn("<w:wordWrap", styles + body)
        self.assertIn('w:w="11906"', body)  # A4 width in twips
        self.assertIn("2026. 10. 1.", body)  # no ISO date on the page
        self.assertNotIn("2026-10-01", body)
        self.assertIn("&lt;표 1&gt;", body)
        header_xml = "".join(_xml(out, n) for n in zipfile.ZipFile(out).namelist() if n.startswith("word/header"))
        self.assertNotIn("예시 데이터", header_xml, "the notice is said once, never as a header chip")

    def test_plain_build_is_a4_pretendard_and_publisher_and_tonality_do_not_mix(self):
        plain = self.build()[0]
        self.assertIn('w:w="11906"', _xml(plain))
        self.assertNotIn("TableGrid", _xml(plain))
        self.assertIn('w:ascii="Pretendard"', _xml(plain, "word/styles.xml"))
        run, _o, _s = _convert(self.tmp, "--tonality", "Report", "--publisher", "elsevier", name="both")
        self.assertNotEqual(run.returncode, 0)
        self.assertRegex(run.stderr + run.stdout, r"(?i)tonality.*publisher|publisher.*tonality")

    def test_publisher_build_keeps_component_attribute_quotes_and_curls_prose_quotes(self):
        source = """---
title: 정비 거점 운영 점검
---

::: callout kind=key title="결정 요청"
남부 거점의 야간 운영을 12월 회의에서 결정해 주십시오.
:::

# 점검 배경

세 거점 가운데 남부를 "우선 점검" 대상으로 정했다.
"""
        for publisher in ("korean-generic", "elsevier"):
            out, _src, _log = self.build("--publisher", publisher, source=source, name=f"pub-{publisher}")
            document = docx.Document(str(out))
            text = "\n".join([p.text for p in document.paragraphs]
                             + [c.text for t in document.tables for r in t.rows for c in r.cells])
            self.assertIn("결정 요청", text, publisher)
            self.assertNotIn(":::", text, publisher)
            self.assertIn("“우선 점검”", text, publisher)

    def test_gate_fails_sentence_headings_and_passes_labels(self):
        layout = _load("docx_layout")
        for text in ("운영 결과가 좋았다", "비용이 줄었음", "Costs fell sharply"):
            self.assertTrue(layout.declarative(text), text)
        for text in ("공용 자전거 정비 거점 1년 운영 결과", "비용과 효과", "Maintenance cost by depot"):
            self.assertFalse(layout.declarative(text), text)
        claim = SOURCE.replace("# 비용과 효과", "# 비용이 크게 줄었다")
        out, src, _ = self.build("--tonality", "Report", source=claim, name="claim")
        gate = subprocess.run([sys.executable, str(SCRIPTS / "docx_gate.py"), str(out), "--source", str(src)],
                              capture_output=True, text=True, env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
        self.assertEqual(gate.returncode, 1)
        self.assertIn("heading.declarative", json.loads(gate.stdout)["failure_reasons"][-1])

    def test_gate_passes_a_restrained_report_with_rendered_pages(self):
        if not shutil.which("soffice"):
            self.skipTest("LibreOffice not installed")
        out, src, _ = self.build("--tonality", "Report")
        gate = subprocess.run([sys.executable, str(SCRIPTS / "docx_gate.py"), str(out), "--source", str(src), "--layout"],
                              capture_output=True, text=True, env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"}, timeout=400)
        report = json.loads(gate.stdout)
        self.assertTrue(report["output"]["rendered"], report["output"].get("render"))
        self.assertEqual(gate.returncode, 0, report["failure_reasons"])

    MEMO = """---
title: New opening hours for the tool crib from 3 November
author: Maintenance planning (example)
organization: Example Plant
date: 2026-10-01
short_title: Tool crib hours
notice: "Example memo: names, dates and numbers are placeholders"
---

::: cover variant=masthead kicker="Internal memo"
To: All shift technicians (example) · From: Maintenance planning · Subject: tool crib hours
:::

# What changes

From Monday 3 November, the tool crib in hall B opens from 06:00 to 22:00 instead of around the clock. Between 22:00 and 06:00, night-shift technicians draw tools from the locked cabinet beside the hall B control room. The cabinet holds the 40 tools that the night shift used most often this year.

We are making this change because the crib now issues about 15 tools a night, down from 70 two years ago. Keeping a storekeeper on duty for so few requests has left the day shift short of one person three times a week.

::: callout kind=warning title="Calibrated tools"
Torque wrenches and gauges with a calibration label never go into the night cabinet. If a night job needs one, call the shift lead, who holds the key to the calibrated drawer.
:::

# What you need to do

1. Book any tool you need for a night job before 22:00 in the maintenance system. The crib will set it aside with your name and job number.
2. Sign each tool out of the night cabinet on the sheet inside the door, with the time and the job number.
3. Return night tools to the crib counter by 08:00. The storekeeper checks them back in and restocks the cabinet.

Technicians on call from home keep using the van kits. Nothing changes for them.

# Numbers behind the decision

Table 1. Night requests at the crib, two years apart (example data)

| Measure (example) | Oct 2024 | Sep 2026 |
| --- | ---: | ---: |
| Tools issued per night | 70 | 15 |
| Distinct tools requested per month | 210 | 52 |
| Day shifts short of a storekeeper per week | 0 | 3 |

The cabinet's 40 tools covered 38 of the 52 distinct tools requested at night last month.

# Questions

Ask maintenance planning on extension 2280 (example) or reply to this memo. We will review the arrangement at the end of January and report back to all shifts.
"""

    def _gate(self, out, src, *flags):
        run = subprocess.run([sys.executable, str(SCRIPTS / "docx_gate.py"), str(out), "--source", str(src), *flags],
                             capture_output=True, text=True, env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"}, timeout=400)
        return json.loads(run.stdout)

    @staticmethod
    def _spacing(xml, text, key):
        """The w:spacing attribute `key` (twips) of the paragraph whose text starts with `text`."""
        at = xml.index(f">{text}")
        start = max(xml.rfind("<w:p>", 0, at), xml.rfind("<w:p ", 0, at))
        ppr = re.search(r"<w:pPr>[\s\S]*?</w:pPr>", xml[start:at])
        found = re.search(rf'<w:spacing\b[^>]*w:{key}="(\d+)"', ppr.group(0) if ppr else "")
        return int(found.group(1)) if found else 0

    def test_memo_that_fits_one_page_stays_on_one_page_with_clear_run_in_head_and_list_end(self):
        if not shutil.which("soffice"):
            self.skipTest("LibreOffice not installed")
        out, src, _ = self.build("--tonality", "Memo", source=self.MEMO, name="memo")
        xml = _xml(out)
        self.assertGreaterEqual(self._spacing(xml, "Calibrated tools", "before"), 240, "a run-in head starts a block")
        self.assertGreaterEqual(self._spacing(xml, "Technicians on call", "before"), 120, "the paragraph after a list stands clear")
        report = self._gate(out, src, "--layout")
        self.assertTrue(report["output"]["rendered"], report["output"].get("render"))
        self.assertEqual(report["output"].get("pages"), 1, report["output"].get("fill"))
        self.assertNotIn("memo.fit", [f["check"] for f in report["output"]["findings"]])
        spill = self.MEMO + ("\nThe new hours apply to contractors' technicians as well; they book tools through their supervisor and "
                             "collect them at the crib counter before 22:00.\n\nWe will post the new hours at the crib window in the week "
                             "before 3 November and update the shift handover sheet on the same day.\n")
        long_out, long_src, _ = self.build("--tonality", "Memo", source=spill, name="memo-long")
        self.assertIn("memo.fit", [f["check"] for f in self._gate(long_out, long_src, "--layout")["output"]["findings"]])

    def test_publisher_text_stands_clear_of_tables_and_captions(self):
        from PIL import Image
        Image.new("RGB", (60, 30), (120, 120, 120)).save(Path(self.tmp) / "fig.png")
        source = """---
title: Wear of a coated drive bearing
author: Example Author A
---

# Results

The coated bearing loses less mass than the plain bearing at every load.

Table 1. Mass loss after 200 h (example values)

| Load (kN) | Plain | Coated |
| ---: | ---: | ---: |
| 4 | 41 | 23 |
| 6 | 50 | 28 |

Across both loads the coated bearing loses about 45 % less mass by 200 h.

![Figure 1. Mass loss over time (synthetic)](fig.png)

The two curves also differ in shape.
"""
        for publisher in ("elsevier", "acs"):
            out, _src, _log = self.build("--publisher", publisher, source=source, name=f"pub-space-{publisher}")
            xml = _xml(out)
            self.assertGreaterEqual(self._spacing(xml, "Across both loads", "before"), 160, f"{publisher}: the paragraph after the table")
            caption = xml[:xml.index("Mass loss after")]
            head = re.search(r'<w:spacing\b[^>]*w:before="(\d+)"', caption[caption.rindex("<w:pPr>"):])
            self.assertGreaterEqual(int(head.group(1)) if head else 0, 160, f"{publisher}: the table caption after text")
            figure = xml[:xml.index("Mass loss over time")]
            tail = re.search(r'<w:spacing\b[^>]*w:after="(\d+)"', figure[figure.rindex("<w:pPr>"):])
            self.assertGreaterEqual(int(tail.group(1)) if tail else 0, 160, f"{publisher}: the figure caption keeps 8 pt under it")

    # ── fix round 3: short lists and headings kept together, tight one-page sources, balanced tables ──

    @staticmethod
    def _ppr(xml, text):
        at = xml.index(text)
        start = max(xml.rfind("<w:p>", 0, at), xml.rfind("<w:p ", 0, at))
        found = re.search(r"<w:pPr>[\s\S]*?</w:pPr>", xml[start:at])
        return found.group(0) if found else ""

    def test_a_short_list_keeps_together_with_its_lead_in_and_a_heading_keeps_its_lead_with_the_table(self):
        source = """---
title: 정비 거점 운영 결과
notice: "예시 데이터: 수치는 가정입니다"
---

# 결론과 다음 단계

정비 거점은 처리 시간과 재정비 비율에서 기대한 효과를 냈다.

다음 단계는 아래와 같다.

1. 2026년 11월: 거점 두 곳 추가 예산 승인
2. 2027년 1월: 주말 두 차례에 걸쳐 작업대 설치
3. 2027년 2월: 야간 당직 인력 1명 추가 배치
4. 2027년 3월: 효과 확인 뒤 다음 거점 설계 착수

이 보고서의 모든 수치는 예시다.

# 시나리오 비교

세 시나리오의 투자비와 회수 기간을 비교했다.

표 1. 확대 시나리오 비교 (예시)

| 시나리오 | 투자비 (억 원) | 회수 기간 (년) |
| --- | ---: | ---: |
| A. 현재 거점 | 1.8 | 6.2 |
| B. 두 곳 추가 | 2.9 | 5.0 |

자료: 예시 추정
"""
        out, _src, _ = self.build("--tonality", "Report", source=source, name="d1")
        xml = _xml(out)
        keep = lambda text: "<w:keepNext/>" in self._ppr(xml, text)  # noqa: E731
        self.assertTrue(keep("다음 단계는 아래와 같다."), "the lead-in keeps with its list")
        for item in ("2026년 11월", "2027년 1월", "2027년 2월"):
            self.assertTrue(keep(item), f"{item} keeps with the next item")
        self.assertFalse(keep("2027년 3월"), "the last item ends the chain")
        self.assertTrue(keep("세 시나리오의 투자비와"), "the heading's lead sentence keeps with the table it opens")

    def test_the_page_checks_find_a_split_list_a_heading_apart_a_near_empty_page_and_a_lone_last_paragraph(self):
        layout = _load("docx_layout")
        frame = {"page_h": 841.9, "top": 70.0, "bottom": 70.0}

        def page(lines, bottom):
            return {"h": 841.9, "lines": lines, "ink": [(y0, y1) for y0, y1, _ in lines] + [(70.0, bottom)], "images": 0,
                    "boxes": [(72.0, y0, 300.0, y1, t) for y0, y1, t in lines], "filled": 0.0}
        blocks = [{"t": "h", "level": 1, "text": "결론", "raw": "결론"}, {"t": "p", "text": "다음단계는아래와같다", "list": False},
                  {"t": "p", "text": "첫째항목입니다", "list": True}, {"t": "p", "text": "둘째항목입니다", "list": True},
                  {"t": "p", "text": "셋째항목입니다", "list": True}, {"t": "h", "level": 1, "text": "비교", "raw": "비교"},
                  {"t": "p", "text": "세시나리오를비교했다", "list": False},
                  {"t": "tbl", "rows": ["시나리오투자비", "A현재거점"], "component": None, "cells": []}]
        info = {"blocks": blocks, "components": {k: 0 for k in layout.KINDS}, "cover": False, "memo": False, "frame": frame}
        p1 = page([(80, 92, "결론"), (100, 112, "다음단계는아래와같다"), (700, 712, "첫째항목입니다")], 712)
        p2 = page([(80, 92, "둘째항목입니다"), (100, 112, "셋째항목입니다"), (740, 752, "비교"), (760, 770, "세시나리오를비교했다")], 770)
        p3 = page([(80, 90, "시나리오투자비"), (95, 105, "A현재거점")], 105)
        found, _ = layout.paged(info, [p1, p2, p3], True, {})
        checks = {f["check"] for f in found}
        for check in ("list.split", "heading.apart", "page.spill"):
            self.assertIn(check, checks)
        found, _ = layout.paged({**info, "blocks": blocks[:1]}, [p1, p1, page([(80, 92, "끝")], 92)], True, {})
        self.assertIn("page.spill", {f["check"] for f in found})
        # A heading whose short section ends a page is not apart from the next section's table.
        other = blocks[:2] + [{"t": "h", "level": 1, "text": "예산", "raw": "예산"}] + blocks[6:]
        q1 = page([(80, 92, "결론"), (100, 112, "다음단계는아래와같다")] + [(130 + 20 * i, 142 + 20 * i, f"본문{i}줄입니다") for i in range(30)], 742)
        q2 = page([(80, 92, "예산"), (100, 112, "세시나리오를비교했다"), (130, 140, "시나리오투자비"), (145, 155, "A현재거점")], 155)
        found, _ = layout.paged({**info, "blocks": other}, [q1, q2], True, {})
        self.assertNotIn("heading.apart", {f["check"] for f in found})

    def test_a_source_of_about_a_page_is_set_tight_in_any_tonality_and_a_section_after_a_ruled_box_draws_no_second_rule(self):
        if not shutil.which("soffice"):
            self.skipTest("LibreOffice not installed")
        out, src, stdout = self.build("--tonality", "Brief", source=self.MEMO, name="brief")
        self.assertIn("tight spacing", stdout)
        self.assertIn('<w:top w:val="nil"/>', self._ppr(_xml(out), "What you need to do"), "no hairline over the section under the box")
        report = self._gate(out, src, "--layout", "--tonality", "Brief")
        self.assertTrue(report["output"]["rendered"], report["output"].get("render"))
        self.assertNotIn("page.spill", [f["check"] for f in report["output"]["findings"] + report["output"].get("advisories", [])])
        self.assertEqual(report["output"].get("pages"), 1, report["output"].get("fill"))

    def test_a_wrapping_text_column_takes_room_negatives_read_with_the_minus_sign_and_an_uncaptioned_table_stands_clear(self):
        source = """---
title: 거점별 주간 정비
notice: "예시 데이터: 수치는 가정입니다"
---

# 거점별 정비

거점마다 정비 건수와 변화를 적었다.

| Depot | Repairs | Change | Share |
| --- | ---: | ---: | ---: |
| Harbour Road | 1,240 | -3 | 21% |
| North Gate Workshop | 980 | -12 | 17% |
| Old Mill | 2,105 | 8 | 36% |

자료: 예시 집계
"""
        out, _src, _ = self.build("--tonality", "Report", source=source, name="d3")
        xml = _xml(out)
        grid = [int(w) / 20 for w in re.findall(r'<w:gridCol w:w="(\d+)"/>', xml)]
        self.assertGreaterEqual(grid[0], 95, f"the depot column holds most of its longest name: {grid}")
        self.assertIn("−3<", xml)
        self.assertIn("−12<", xml)
        self.assertNotIn(">-3<", xml)
        self.assertGreaterEqual(self._spacing(xml, "거점마다 정비 건수와", "after"), 120, "the sentence over an uncaptioned table keeps 6 pt")

    def test_under_a_publisher_a_paragraph_after_a_list_and_after_a_table_stands_clear_and_a_short_list_keeps_together(self):
        source = """---
title: A five-line inspection template
author: Example Author A
---

# Results

We propose that every short inspection report states five lines.

1. Bearing temperature at the end of the run.
2. Time to steady temperature.
3. Peak vibration.

Lines 3 and 4 together let a reader compute the margin.

Table 1. Margin (example)

| Speed (m/s) | Margin (°C) |
| ---: | ---: |
| 1.0 | 8.8 |
| 2.0 | -0.7 |

The negative margin at 2.0 m/s does not make the bearing unusable.
"""
        for publisher in ("elsevier", "acs"):
            out, _src, _ = self.build("--publisher", publisher, source=source, name=f"d4-{publisher}")
            xml = _xml(out)
            self.assertGreaterEqual(self._spacing(xml, "Lines 3 and 4", "before"), 120, publisher)
            self.assertGreaterEqual(self._spacing(xml, "The negative margin", "before"), 180, publisher)
            self.assertTrue(all("<w:keepNext/>" in self._ppr(xml, t) for t in ("Bearing temperature at", "Time to steady")), publisher)

    def test_the_gate_fails_a_notice_that_joins_its_label_and_line_with_a_spaced_dash(self):
        dashed = "---\ntitle: 운영 메모\ntonality: Memo\nnotice: 예시 안내문 — 일정은 가정입니다\n---\n\n# 일정\n\n다음 주 월요일부터 바뀐다.\n"
        out, src, _ = self.build("--tonality", "Memo", source=dashed, name="dash")
        checks = [f"{f['severity']}:{f['check']}" for f in self._gate(out, src)["output"]["findings"]]
        self.assertIn("FAIL:notice.dash", checks)
        colon = dashed.replace("notice: 예시 안내문 — 일정은 가정입니다", 'notice: "예시 안내문: 일정은 가정입니다"')
        out, src, _ = self.build("--tonality", "Memo", source=colon, name="colon")
        self.assertNotIn("notice.dash", [f["check"] for f in self._gate(out, src)["output"]["findings"]])


    def test_a_two_column_body_spans_a_table_source_lets_a_column_list_break_once_and_closes_with_a_continuous_break(self):
        source = """---
title: Bearing wear on two conveyor lines
author: Example Author A
notice: "Synthetic example: data are placeholders"
---

# Method

Each line was inspected weekly for twelve weeks.

Table 1. Inspections by line and shift (example values)

| Line | Shift | Bearings | Worn (%) |
| --- | --- | ---: | ---: |
| North | Day | 240 | 3.1 |
| South | Night | 180 | 5.4 |

Source: synthetic data, twelve weeks

## Inspection

Each bearing was checked for play and heat.

Every short wear report states five lines:

1. Bearings checked per week.
2. Share found worn.
3. Longest run without a stop.
4. Spare bearings on hand.
5. Number of weeks logged.

Lines 3 and 4 together give the spare margin.
"""
        out, _src, _ = self.build("--tonality", "Journal", source=source, name="j1")
        xml = _xml(out)
        table = xml.index("Inspections by line and shift")
        resume = xml.index('w:num="1"', table)
        self.assertLess(xml.index(">Source: synthetic data"), resume, "the source line stands under the page-wide table")
        keep = lambda text: "<w:keepNext/>" in self._ppr(xml, text)  # noqa: E731
        for text in ("Every short wear report", "Bearings checked per week", "Spare bearings on hand"):
            self.assertTrue(keep(text), f"{text} keeps with the next paragraph")
        for text in ("Share found worn", "Longest run without a stop"):
            self.assertFalse(keep(text), f"the column list may break after {text}")
        sects = list(re.finditer(r"<w:sectPr\b[\s\S]*?</w:sectPr>", xml))
        self.assertNotIn('w:num="2"', sects[-1].group(0), "the document closes on a one-column section")
        self.assertIn('w:num="2"', sects[-2].group(0))
        self.assertGreater(sects[-2].start(), xml.index("Lines 3 and 4 together"))
        self.assertIn('<w:widowControl w:val="0"/>', self._ppr(xml, "Lines 3 and 4 together"), "the last paragraph may break across the balanced columns")

    def test_a_short_closing_section_keeps_whole_and_one_with_a_table_does_not(self):
        head = "---\ntitle: 정비 거점 운영 결과\nnotice: \"예시 데이터: 수치는 가정입니다\"\n---\n\n# 배경\n\n정비 거점을 두 곳에서 여섯 달 운영했고, 처리 시간과 재정비 비율, 부품 대기를 주마다 기록해 계획과 비교했다.\n\n"
        tail = ("# 결론과 다음 단계\n\n정비 거점은 처리 시간과 재정비 비율에서 기대한 효과를 냈다. 부품 대기가 남은 원인은 거점 자체가 아니라 "
                "발주 시점과 야간 당직에 있었고, 두 가지 모두 고칠 수 있다.\n\n다음 단계는 아래와 같다.\n\n1. 2026년 11월: 거점 두 곳 추가 예산 승인\n"
                "2. 2027년 1월: 작업대 설치\n3. 2027년 2월: 야간 당직 인력 추가 배치\n\n이 보고서의 모든 수치는 예시다.\n")
        out, _src, _ = self.build("--tonality", "Report", source=head + tail, name="j2")
        xml = _xml(out)
        self.assertIn("<w:keepNext/>", self._ppr(xml, "정비 거점은 처리 시간과"), "the closing paragraph keeps with its list")
        self.assertNotIn("<w:keepNext/>", self._ppr(xml, "정비 거점을 두 곳에서"), "an earlier section is untouched")
        out, _src, _ = self.build("--tonality", "Report", source=head + tail + "\n| 단계 | 시기 |\n| --- | --- |\n| 승인 | 11월 |\n", name="j2t")
        self.assertNotIn("<w:keepNext/>", self._ppr(_xml(out), "정비 거점은 처리 시간과"), "a closing section with a table moves no paragraph")

    def test_the_page_checks_find_a_heading_at_a_column_foot_and_columns_that_do_not_balance(self):
        layout = _load("docx_layout")
        frame = {"page_h": 841.9, "top": 70.0, "bottom": 70.0}

        def page(lines):
            return {"h": 841.9, "w": 595.3, "lines": sorted((y0, y1, t) for _, y0, y1, t in lines), "ink": [(y0, y1) for _, y0, y1, _ in lines],
                    "images": 0, "boxes": [(x0, y0, x0 + 230.0, y1, t) for x0, y0, y1, t in lines], "filled": 0.0}

        def col(x, y0, y1, tag):
            return [(x, y, y + 11, f"{tag}{k}줄") for k, y in enumerate(range(int(y0), int(y1), 14))]

        def blocks(*items):
            return [dict({"t": t, "text": x, "list": lst, "cols": 2}, **({"level": 1, "raw": x} if t == "h" else {})) for t, x, lst in items]

        def checks(b, pages):
            info = {"blocks": b, "components": {k: 0 for k in layout.KINDS}, "cover": False, "memo": False, "frame": frame}
            return {f["check"] for f in layout.paged(info, pages, True, {})[0]}
        heads = blocks(("h", "배경", False), ("p", "왼쪽0줄", False), ("h", "방법", False), ("p", "다음0줄", False))
        first = page([(72, 80, 92, "배경")] + col(72, 102, 760, "왼쪽") + col(310, 80, 740, "오른쪽") + [(310, 750, 762, "방법")])
        self.assertIn("heading.column", checks(heads, [first, page(col(72, 80, 420, "다음") + col(310, 80, 410, "끝"))]))
        self.assertIn("columns.balance", checks(heads[:2], [first, page(col(72, 80, 500, "다음"))]))
        listed = blocks(("h", "배경", False), ("p", "왼쪽0줄", False), *[("p", f"{n}항목", True) for n in "가나다라마"])
        p1 = page([(72, 80, 92, "배경")] + col(72, 102, 730, "왼쪽") + [(72, 740, 751, "가항목"), (72, 754, 765, "나항목")] + col(310, 80, 765, "오른쪽"))
        p2 = page([(72, 80, 91, "다항목"), (72, 94, 105, "라항목"), (72, 108, 119, "마항목")] + col(72, 124, 420, "다음") + col(310, 80, 420, "끝"))
        self.assertFalse(checks(listed, [p1, p2]) & {"columns.balance", "list.split", "heading.column"})
        short = page([(72, 80, 92, "배경")] + col(72, 100, 520, "왼쪽") + col(310, 80, 765, "오른쪽"))
        self.assertIn("columns.balance", checks(listed[:2], [short, p2]))


if __name__ == "__main__":
    unittest.main()
